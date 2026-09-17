from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = '/Users/F7905607/Dropbox/Projects/hoe/app/Quicken Tree Loyalty App Features.docx'

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tc_pr.append(shd)

def borders(cell, color='D9D9D9'):
    tc_pr = cell._tc.get_or_add_tcPr(); tc_borders = tc_pr.first_child_found_in('w:tcBorders')
    if tc_borders is None:
        tc_borders = OxmlElement('w:tcBorders'); tc_pr.append(tc_borders)
    for edge in ('top','left','bottom','right'):
        tag = 'w:' + edge; element = tc_borders.find(qn(tag))
        if element is None: element = OxmlElement(tag); tc_borders.append(element)
        element.set(qn('w:val'),'single'); element.set(qn('w:sz'),'6'); element.set(qn('w:color'),color)

def set_cell_text(cell, text, bold=False, color='111111', size=10.5):
    cell.text = ''
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2)
    r = p.add_run(text); r.bold = bold; r.font.size = Pt(size); r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; borders(cell)

def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr(); header = OxmlElement('w:tblHeader'); header.set(qn('w:val'), 'true'); tr_pr.append(header)

def keep_row_together(row):
    tr_pr = row._tr.get_or_add_trPr(); cant_split = OxmlElement('w:cantSplit'); tr_pr.append(cant_split)

def remove_title_border(paragraph):
    p_pr = paragraph._p.get_or_add_pPr(); p_bdr = p_pr.find(qn('w:pBdr'))
    if p_bdr is not None: p_pr.remove(p_bdr)

doc = Document()
sec = doc.sections[0]; sec.top_margin = Inches(1.0); sec.bottom_margin = Inches(.7); sec.left_margin = Inches(.75); sec.right_margin = Inches(.75)
styles = doc.styles
styles['Normal'].font.name = 'Source Sans Pro'; styles['Normal'].font.size = Pt(10.5)
styles['Normal'].paragraph_format.space_after = Pt(7)
for name, size in [('Title',24),('Heading 1',15),('Heading 2',12)]:
    styles[name].font.name = 'Source Sans Pro'; styles[name].font.size = Pt(size); styles[name].font.bold = True; styles[name].font.color.rgb = RGBColor(0,0,0)
title_style_pr = styles['Title'].element.get_or_add_pPr()
title_style_border = title_style_pr.find(qn('w:pBdr'))
if title_style_border is not None: title_style_pr.remove(title_style_border)

title = doc.add_paragraph(style='Title'); title.alignment = WD_ALIGN_PARAGRAPH.LEFT; title.add_run('Quicken Tree Loyalty App Features'); remove_title_border(title)
sub = doc.add_paragraph(); sub.add_run('Product brief for bookings loyalty and footfall growth').bold = True
sub.runs[0].font.color.rgb = RGBColor(197,22,43)
doc.add_paragraph('The app should make booking at The Quicken Tree simple across breakfast lunch and evening service while giving guests a clear reason to choose quieter times. The central mechanism is a loyalty programme that rewards repeat visits and uses targeted double points to help build footfall when the restaurant has capacity.')

doc.add_heading('Core customer experience', level=1)
for heading, body in [
    ('Book a table', 'Guests select their party size date and time, with availability shown across breakfast lunch and evening service. The booking journey should confirm the points offer attached to the selected slot before checkout.'),
    ('Earn and redeem points', 'Members see their live balance, the next reward threshold and a practical choice of rewards such as a complimentary coffee, a drink on the house or dining credit. Points are added after a completed visit.'),
    ('Member profile', 'The profile keeps upcoming bookings, dietary and occasion preferences, favourite dishes and reward history in one place. This helps staff recognise returning guests and makes future bookings quicker.'),
    ('What is on', 'A lightweight events area promotes supper clubs, seasonal menus, live nights and family activities. Members can save an event or book a table directly from the listing.')
]:
    p = doc.add_paragraph(); p.add_run(heading + '. ').bold = True; p.add_run(body)

doc.add_heading('Booking availability', level=1)
doc.add_paragraph('The app must represent the venue as an all day destination rather than an evenings only restaurant. The proposed public booking windows are below. Availability inside each window should remain controlled by the live table plan and service capacity.')
table = doc.add_table(rows=1, cols=3); table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.style = 'Table Grid'
headers = ['Day', 'Booking window', 'Services supported']
for cell, text in zip(table.rows[0].cells, headers): shade(cell,'111111'); set_cell_text(cell,text,True,'FFFFFF')
repeat_header(table.rows[0])
for row in [('Monday to Friday','7:30 AM to 8:30 PM','Breakfast lunch and evening dining'),('Saturday','7:30 AM to 8:30 PM','Breakfast lunch and evening dining'),('Sunday','7:30 AM to 5:30 PM','Breakfast lunch and Sunday dining')]:
    cells = table.add_row().cells
    keep_row_together(table.rows[-1])
    for cell, text in zip(cells,row): set_cell_text(cell,text)

doc.add_heading('Footfall and loyalty mechanics', level=1)
doc.add_paragraph('Double points should be used as a targeted incentive, not as a permanent discount. The app can apply the multiplier automatically when a guest books an eligible slot and show the expected points before they confirm.')
table = doc.add_table(rows=1, cols=3); table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.style = 'Table Grid'
for cell, text in zip(table.rows[0].cells,['Feature','How it works','Footfall benefit']): shade(cell,'C5162B'); set_cell_text(cell,text,True,'FFFFFF')
repeat_header(table.rows[0])
for row in [
    ('Double points slots','Show a Double Points badge on selected booking times, such as quieter weekday breakfast, lunch or early evening periods.','Moves demand into lower occupancy periods without reducing menu price.'),
    ('Breakfast and lunch boosts','Offer a short booking bonus or coffee reward for repeat daytime visits.','Makes the app relevant beyond dinner and encourages more frequent visits.'),
    ('Return visit nudges','After a completed visit, send an in app reminder when a member is close to a reward.','Creates a clear reason to return rather than letting points go unnoticed.'),
    ('Event access','Give members early booking access or bonus points for supper clubs and special menus.','Builds demand ahead of events and gives the loyalty scheme a distinct benefit.'),
    ('Birthday and occasion rewards','Let customers add an occasion to a booking and unlock a small member treat.','Supports group bookings and celebration occasions with a personal prompt.')
]:
    cells = table.add_row().cells
    keep_row_together(table.rows[-1])
    for cell, text in zip(cells,row): set_cell_text(cell,text, size=9.5)

doc.add_heading('Operational safeguards', level=1)
for text in [
    'Restaurant managers choose which dates and times receive double points, with a simple start and end date for each offer.',
    'The booking system should limit reward redemptions and double points offers where service is already near capacity.',
    'Staff need a clear check in view showing a guest name, booking, loyalty status and any redeemed reward.',
    'The team should review booking conversion, average covers by service period, repeat visit rate and redeemed rewards monthly.'
]: doc.add_paragraph(text, style='List Bullet')

doc.add_heading('Recommended launch sequence', level=1)
doc.add_paragraph('Launch bookings, points balance and three simple rewards first. Next, introduce double points on selected low demand booking slots and measure whether those periods gain covers without moving existing high demand bookings. Once the team is comfortable, add events, occasion prompts and more personalised offers.')

doc.save(OUT)
print(OUT)
